"""Create an editable review DOCX from the project's supported LaTeX subset.

Retains original protected front matter and uses isolated styles for the new body.
Does not claim Word pagination has been verified when Word is unavailable.
"""
from __future__ import annotations
import copy
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT/'docs'/'report'
SOURCE = REPORT/'source'/'original-report.docx'
OUTPUT = ROOT/'output'/'report'/'NEXOS_Professional_Report_Editable.docx'
QA = ROOT/'tmp'/'editable-report'
QA.mkdir(parents=True,exist_ok=True)
IMG = REPORT/'images'
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def el(tag, **attrs):
    e=OxmlElement('w:'+tag)
    for k,v in attrs.items(): e.set(qn('w:'+k),str(v))
    return e

def group(s,pos):
    assert s[pos]=='{',(pos,s[pos:pos+50])
    depth=1; j=pos+1
    while depth and j<len(s):
        if s[j]=='{' and s[j-1]!='\\': depth+=1
        if s[j]=='}' and s[j-1]!='\\': depth-=1
        j+=1
    return s[pos+1:j-1],j

BIB={key:str(i+1) for i,key in enumerate(re.findall(r'\\bibitem\{([^}]+)\}',(REPORT/'references.tex').read_text()))}
REF={'ch:evaluation':'11'}

def plain(s):
    s=s.replace('~',' ').replace('``','“').replace("''",'”')
    s=re.sub(r'\\cite\{([^}]+)\}',lambda m:'['+', '.join(BIB[k] for k in m[1].split(','))+']',s)
    s=re.sub(r'\\ref\{([^}]+)\}',lambda m:REF.get(m[1],m[1]),s)
    s=s.replace('\\_','_').replace('\\&','&').replace('\\%','%')
    for cmd in ['textbf','emph','texttt','textit','url','operatorname','text']:
        while '\\'+cmd+'{' in s:
            start=s.index('\\'+cmd+'{'); brace=start+len(cmd)+1
            val,end=group(s,brace)
            s=s[:start]+val+s[end:]
    s=s.replace('\\rightarrow','→').replace('\\land',' ∧ ').replace('\\Vert',' ∥ ')
    s=s.replace('\\operatorname','').replace('\\quad',' ').replace('\\par','')
    s=s.replace('\\textwidth','').replace('\\noindent','').replace('\\hfill','    ')
    s=s.replace('$','').replace('\\(','').replace('\\)','')
    return re.sub(r'\s+',' ',s).strip()

def runtext(p,s,bold=False,italic=False,mono=False):
    # Preserve common inline emphasis while all body text remains editable.
    pattern=re.compile(r'\\(textbf|emph|textit|texttt)\{')
    pos=0
    for_loop=0
    while True:
        m=pattern.search(s,pos)
        if not m: break
        if m.start()>pos:
            r=p.add_run(plain(s[pos:m.start()]) + (' ' if s[m.start()-1:m.start()].isspace() else ''))
            r.bold=bold;r.italic=italic
            if mono:r.font.name='Consolas'
        val,end=group(s,m.end()-1)
        r=p.add_run(plain(val));r.bold=bold or m[1]=='textbf';r.italic=italic or m[1] in ('emph','textit')
        if mono or m[1]=='texttt':r.font.name='Consolas';r.font.size=Pt(10.5)
        if end<len(s) and s[end].isspace():p.add_run(' ')
        pos=end
        for_loop+=1
        if for_loop>100:raise ValueError('Inline parser loop')
    if pos<len(s):
        r=p.add_run(plain(s[pos:]));r.bold=bold;r.italic=italic
        if mono:r.font.name='Consolas'

def font(size,bold=False):
    return ImageFont.truetype('C:/Windows/Fonts/'+('arialbd.ttf' if bold else 'arial.ttf'),size)

def diagram(name,title,boxes,arrows,height=1050):
    im=Image.new('RGB',(1600,height),'white'); d=ImageDraw.Draw(im)
    d.text((55,35),title,font=font(38,True),fill='#172B40')
    for a,b in arrows:
        x1,y1,x2,y2=boxes[a][0];u1,v1,u2,v2=boxes[b][0]
        if v1>=y2:
            p=((x1+x2)//2,y2);q=((u1+u2)//2,v1)
            d.line((p,q),fill='#1E4D70',width=4)
            d.polygon([(q[0],q[1]),(q[0]-9,q[1]-16),(q[0]+9,q[1]-16)],fill='#1E4D70')
        else:
            p=(x2,(y1+y2)//2);q=(u1,(v1+v2)//2)
            d.line((p,q),fill='#1E4D70',width=4)
            d.polygon([q,(q[0]-16,q[1]-9),(q[0]-16,q[1]+9)],fill='#1E4D70')
    for rect,lines in boxes:
        d.rounded_rectangle(rect,12,fill='#EDF2F6',outline='#1E4D70',width=3)
        total=len(lines)*39;y=(rect[1]+rect[3]-total)//2
        for idx,line in enumerate(lines):
            f=font(29,idx==0);w=d.textlength(line,font=f)
            d.text(((rect[0]+rect[2]-w)/2,y),line,font=f,fill='#172B40');y+=39
    if name=='word-data-model.png':
        d.line([(130,220),(75,220),(75,800),(130,800)],fill='#1E4D70',width=4)
        d.polygon([(130,800),(114,791),(114,809)],fill='#1E4D70')
        for pos in [(735,180),(440,345),(1210,345),(1210,645)]:
            d.text(pos,'1 : many',font=font(22),fill='#4D6072')
    d.text((55,height-55),'NEXOS proposed design',font=font(23),fill='#4D6072')
    im.save(IMG/name)

diagram('word-deployment.png','Deployment inside a virtual machine',[
((100,120,1500,240),['Windows development host','VMware Workstation']),
((100,300,1500,420),['NEXOS guest operating system','Ubuntu, GNOME, Linux kernel and virtual hardware']),
((100,510,500,660),['Qt desktop','User session']),
((600,510,1000,660),['nexosd','Task engine']),
((1100,510,1500,660),['Ollama','Local inference']),
((330,760,1270,900),['Persistent guest disk','SQLite, task artifacts and model files'])],[(0,1),(2,3),(3,4),(3,5)])
diagram('word-dataflow.png','Data flow and permission boundary',[
((330,120,1270,240),['User request and explicit resource grants']),
((330,300,1270,420),['Planner and schema validation','Project data and model output are untrusted']),
((330,480,1270,600),['Deterministic permission boundary']),
((330,660,1270,780),['Isolated tools or authorized API connector']),
((330,840,1270,960),['Evidence checks and durable result'])],[(0,1),(1,2),(2,3),(3,4)],1080)
diagram('word-data-model.png','Proposed logical data model',[
((130,130,700,310),['TASK','id, goal, status, budget']),
((900,130,1470,310),['STEP','id, task id, tool, dependencies']),
((130,420,700,600),['EVENT','sequence, timestamp, state']),
((900,420,1470,600),['ATTEMPT','step id, outcome, timing']),
((130,710,700,890),['GRANT','action digest, expiry']),
((900,710,1470,890),['ARTIFACT','path, digest, provenance'])],[(0,1),(0,2),(1,3),(3,5)])
diagram('word-boot-qualification.png','Boot and installation qualification',[
((220,115,1380,235),['Attach candidate ISO to a clean virtual machine']),
((220,285,1380,405),['Boot the live session and check readiness']),
((220,455,1380,575),['Install onto an empty virtual disk']),
((220,625,1380,745),['Eject the ISO and reboot the installed system']),
((220,795,1380,915),['Run a task and verify persistent history']),
((220,965,1380,1085),['Record evidence and release limitations'])],[(0,1),(1,2),(2,3),(3,4),(4,5)],1200)

doc=Document(SOURCE)
original_elements=[copy.deepcopy(e) for e in list(doc._element.body)[:74]]
assert 'DECLARATION' in ''.join(list(doc._element.body)[74].itertext())
for e in list(doc._element.body)[74:]:
    if e.tag!=qn('w:sectPr'):doc._element.body.remove(e)

def style(name,size=12,bold=False,spacing=1.5,outline=None,page=False):
    st=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
    rf=st.element.get_or_add_rPr().find(qn('w:rFonts'))
    for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')
    pf=st.paragraph_format;pf.line_spacing=spacing;pf.space_after=Pt(6);pf.widow_control=True
    if outline is not None:
        st.element.get_or_add_pPr().append(el('outlineLvl',val=outline))
        pf.keep_with_next=True;pf.space_before=Pt(15);pf.space_after=Pt(8)
    if page:pf.page_break_before=True
    return st

style('NEXOS Body')
style('NEXOS Chapter',18,True,1.15,0,True)
style('NEXOS Section',14,True,1.15,1)
style('NEXOS Subsection',12,True,1.15,2)
style('NEXOS Appendix',18,True,1.15,0,True)
style('NEXOS App Section',14,True,1.15,1)
style('NEXOS App Subsection',12,True,1.15,2)
style('NEXOS Front',18,True,1.15,0,True)
style('NEXOS Contents Title',18,True,1.15,None,True)
style('NEXOS Contents',11,False,1.08)
style('NEXOS Caption',10.5,False,1.1)
style('NEXOS Table',10.5,False,1.1)
style('NEXOS Code',9,False,1.0)
style('NEXOS Reference',10.5,False,1.12)
style('NEXOS Footer',9,False,1.0)
style('NEXOS List',12,False,1.35)
doc.styles['NEXOS Code'].font.name='Consolas'
doc.styles['NEXOS Code'].paragraph_format.space_after=Pt(0)
doc.styles['NEXOS Body'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY

def numbering(nid,letter=False):
    abstract=el('abstractNum',abstractNumId=nid)
    abstract.append(el('multiLevelType',val='multilevel'))
    for level in range(3):
        lvl=el('lvl',ilvl=level);lvl.append(el('start',val=1))
        lvl.append(el('numFmt',val='upperLetter' if letter and level==0 else 'decimal'))
        lvl.append(el('lvlText',val='.'.join('%'+str(j+1) for j in range(level+1))))
        lvl.append(el('lvlJc',val='left'));lvl.append(el('suff',val='space'))
        abstract.append(lvl)
    numbering_root=doc.part.numbering_part.element
    first_num=numbering_root.find(qn('w:num'))
    if first_num is not None:first_num.addprevious(abstract)
    else:numbering_root.append(abstract)
    num=el('num',numId=nid);num.append(el('abstractNumId',val=nid));doc.part.numbering_part.element.append(num)
numbering(900);numbering(901,True)

section=doc.add_section(WD_SECTION_START.NEW_PAGE)
boundary=doc.paragraphs[-1]
boundary.paragraph_format.space_before=Pt(0);boundary.paragraph_format.space_after=Pt(0)
boundary.paragraph_format.line_spacing=Pt(1)
section.page_width=Mm(210);section.page_height=Mm(297)
section.left_margin=Mm(32);section.right_margin=Mm(25);section.top_margin=Mm(25);section.bottom_margin=Mm(25)
section.header_distance=Mm(12);section.footer_distance=Mm(12)
section.header.is_linked_to_previous=False;section.footer.is_linked_to_previous=False
section.different_first_page_header_footer=False
def set_page_numbering(section,fmt):
    for old in section._sectPr.findall(qn('w:pgNumType')):section._sectPr.remove(old)
    page_num=el('pgNumType',fmt=fmt,start=1)
    before=section._sectPr.find(qn('w:cols'))
    if before is None:before=section._sectPr.find(qn('w:docGrid'))
    if before is not None:before.addprevious(page_num)
    else:section._sectPr.append(page_num)
set_page_numbering(section,'lowerRoman')

def field(p,instruction,value=''):
    r=p.add_run();r._r.append(el('fldChar',fldCharType='begin'))
    r=p.add_run();it=el('instrText');it.set(qn('xml:space'),'preserve');it.text=' '+instruction+' ';r._r.append(it)
    p.add_run()._r.append(el('fldChar',fldCharType='separate'))
    p.add_run(value)
    p.add_run()._r.append(el('fldChar',fldCharType='end'))

p=section.footer.paragraphs[0];p.style='NEXOS Footer';p.alignment=WD_ALIGN_PARAGRAPH.CENTER
field(p,'PAGE','i')
hp=section.header.paragraphs[0];hp.style='NEXOS Footer';hp.add_run('NEXOS  |  Major Project I')

headings=[];figures=[];tables=[];bookmark_id=500
chapter=0;secno=0;subno=0;appendix=False

def bookmark(p,name):
    global bookmark_id
    bookmark_id+=1
    start=el('bookmarkStart',id=bookmark_id,name=name);end=el('bookmarkEnd',id=bookmark_id)
    p._p.insert(1 if p._p.find(qn('w:pPr')) is not None else 0,start);p._p.append(end)

def heading(text,level=0,front=False):
    global chapter,secno,subno
    label=plain(text)
    if front:
        sty='NEXOS Front';shown=label
    else:
        if level==0:chapter+=1;secno=0;subno=0
        elif level==1:secno+=1;subno=0
        else:subno+=1
        c=chr(64+chapter) if appendix else str(chapter)
        prefix=c+('.'+str(secno) if level>=1 else '')+('.'+str(subno) if level==2 else '')
        shown=prefix+'  '+label
        sty=(['NEXOS Appendix','NEXOS App Section','NEXOS App Subsection'] if appendix else ['NEXOS Chapter','NEXOS Section','NEXOS Subsection'])[level]
    p=doc.add_paragraph(style=sty);p.add_run(label)
    if not front:
        np=el('numPr');np.append(el('ilvl',val=level));np.append(el('numId',val=901 if appendix else 900));p._p.get_or_add_pPr().append(np)
    name='nx_heading_'+str(bookmark_id+1);bookmark(p,name)
    if level<2:headings.append((shown,name,level))
    return p

def para(text,sty='NEXOS Body'):
    if not text.strip():return
    p=doc.add_paragraph(style=sty);runtext(p,text);return p

def caption(kind,text):
    entries=figures if kind=='Figure' else tables
    n=len(entries)+1;p=doc.add_paragraph(style='NEXOS Caption')
    p.add_run(kind+' ').bold=True;field(p,'SEQ '+kind+' \\* ARABIC',str(n));p.add_run('. '+plain(text))
    name='nx_'+kind.lower()+'_'+str(n);bookmark(p,name)
    entries.append((f'{kind} {n}. '+plain(text),name,0))
    p.paragraph_format.keep_with_next=kind=='Table'
    return n

def add_figure(path,cap,label=None):
    p=doc.add_paragraph(style='NEXOS Body');p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(3)
    im=Image.open(IMG/path);w,h=im.size
    width=min(153,155*w/h)
    p.add_run().add_picture(str(IMG/path),width=Mm(width))
    num=caption('Figure',cap)
    if label:REF[label]=str(num)

def table(rows,cap=None):
    if cap:caption('Table',cap)
    if not rows:return
    cols=len(rows[0]);assert all(len(r)==cols for r in rows),(cap,rows)
    widths=([31,122] if cols==2 else [25,76,52] if cols==3 else [153/cols]*cols)
    t=doc.add_table(rows=0,cols=cols);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    pr=t._tbl.tblPr
    borders=el('tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:borders.append(el(side,val='single',sz=4,color='D9D9D9'))
    pr.append(borders)
    mar=el('tblCellMar')
    for side in ['top','bottom']:mar.append(el(side,w=85,type='dxa'))
    for side in ['left','right']:mar.append(el(side,w=100,type='dxa'))
    pr.append(mar)
    for j,w in enumerate(widths):t.columns[j].width=Mm(w)
    for idx,row in enumerate(rows):
        cells=t.add_row().cells
        for j,(c,content) in enumerate(zip(cells,row)):
            c.width=Mm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.style='NEXOS Table';runtext(p,content)
            p.paragraph_format.space_after=Pt(0)
            if idx==0:
                c._tc.get_or_add_tcPr().append(el('shd',fill='17324D'))
                for r in p.runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
            elif idx%2==0:c._tc.get_or_add_tcPr().append(el('shd',fill='F3F6F8'))
        rp=t.rows[-1]._tr.get_or_add_trPr();rp.append(el('cantSplit'))
        if idx==0:rp.append(el('tblHeader'))
    doc.add_paragraph(style='NEXOS Caption').paragraph_format.space_after=Pt(3)

def read_table(block):
    cap=None
    m=re.search(r'\\caption\{',block)
    if m:cap,_=group(block,m.end()-1)
    if '\\begin{tabularx}' in block:
        a=block.index('\\begin{tabularx}')+len('\\begin{tabularx}')
        _,a=group(block,a);_,a=group(block,a)
        body=block[a:block.index('\\end{tabularx}')]
    else:
        a=block.index('\\begin{longtable}')+len('\\begin{longtable}')
        _,a=group(block,a);body=block[a:block.index('\\end{longtable}')]
    if '\\endfirsthead' in body:
        header=body[:body.index('\\endfirsthead')]
        tail=body[body.index('\\endhead')+len('\\endhead'):]
        body=header+tail
    body=re.sub(r'\\caption\{[^}]*\}(?:\\label\{[^}]*\})?\\\\','',body)
    body=re.sub(r'\\(?:toprule|midrule|bottomrule|endhead|endfirsthead|centering|small)','',body)
    body=re.sub(r'\\label\{[^}]*\}','',body)
    rows=[]
    for row in re.split(r'\\\\(?:\[[^\]]*\])?',body):
        row=row.strip()
        if row:rows.append([c.strip() for c in re.split(r'(?<!\\)&',row)])
    table(rows,cap)

# Bind cross references before parsing paragraphs that precede their figures.
REF.update({'fig:overview':'1','fig:architecture':'2','fig:deployment':'3','fig:dfd':'4','fig:taskflow':'5','fig:data':'6','fig:ui-task':'7','fig:ui-approval':'8','fig:iso':'9','fig:boot':'10'})
TIKZ={'fig:deployment':'word-deployment.png','fig:dfd':'word-dataflow.png','fig:data':'word-data-model.png','fig:boot':'word-boot-qualification.png'}

def parse(text):
    lines=text.splitlines();i=0;buf=[]
    def flush():
        if buf:para(' '.join(buf));buf.clear()
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:flush();continue
        if line.startswith('%'):continue
        m=re.match(r'\\(chapter|section|subsection)\{',line)
        if m:
            flush();title,_=group(line,m.end()-1);heading(title,{'chapter':0,'section':1,'subsection':2}[m[1]]);continue
        if line.startswith('\\reportfigure'):
            flush();args=[];pos=len('\\reportfigure')
            for _ in range(4):v,pos=group(line,pos);args.append(v)
            add_figure(args[0],args[2],args[3]);continue
        m=re.match(r'\\begin\{(table|longtable|figure|lstlisting|equation|enumerate|itemize)\}',line)
        if m:
            flush();env=m[1];parts=[line]
            while '\\end{'+env+'}' not in parts[-1]:
                if i>=len(lines):raise ValueError('Unclosed '+env)
                parts.append(lines[i]);i+=1
            block='\n'.join(parts)
            if env in ('table','longtable'):read_table(block)
            elif env=='figure':
                cm=re.search(r'\\caption\{',block);cap,_=group(block,cm.end()-1)
                label=re.search(r'\\label\{([^}]+)\}',block)[1]
                add_figure(TIKZ[label],cap,label)
            elif env=='lstlisting':
                cm=re.search(r'caption=\{([^}]+)\}',parts[0])
                if cm:
                    p=para(cm[1],'NEXOS Caption');p.paragraph_format.keep_with_next=True
                for n,code in enumerate(parts[1:-1]):
                    p=doc.add_paragraph(style='NEXOS Code');p.add_run(code)
                    p.paragraph_format.keep_with_next=n<len(parts[1:-1])-1
                doc.add_paragraph(style='NEXOS Caption')
            elif env=='equation':
                if 'VSR' in block:formula='VSR = tasks satisfying all required postconditions / attempted tasks'
                else:formula='allow(a, C) = registered(a) ∧ valid(a) ∧ inScope(a, C)\n                         ∧ policyOK(a) ∧ grantValid(a)'
                p=doc.add_paragraph(style='NEXOS Body');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(formula).italic=True
            else:
                contents=re.sub(r'\\(?:begin|end)\{(?:enumerate|itemize)\}','',block)
                items=re.split(r'\\item\s+',contents)[1:]
                for n,item in enumerate(items,1):
                    p=doc.add_paragraph(style='NEXOS List');p.paragraph_format.left_indent=Mm(6);p.paragraph_format.first_line_indent=Mm(-6)
                    p.add_run(f'{n}. ' if env=='enumerate' else '• ');runtext(p,item)
            continue
        if line.startswith('\\label') or line in ('\\clearpage','\\vfill'):continue
        buf.append(line)
    flush()

# Front matter retains all substantive prose from the approved LaTeX source.
front=(REPORT/'chapters'/'frontmatter.tex').read_text(encoding='utf-8')
for title,body in re.findall(r'\\frontchapter\{([^}]+)\}(.*?)(?=\\frontchapter|\Z)',front,re.S):
    p=heading(title,front=True)
    if title=='Declaration':p.paragraph_format.page_break_before=False
    if title=='Abbreviations and Notation':read_table(body[body.index('\\begin{longtable}'):]);continue
    prose=re.split(r'\\(?:vspace|begin\{tabularx\}|begin\{flushright\})',body)[0]
    prose=prose.replace('\\clearpage','')
    for paragraph in re.split(r'\n\s*\n',prose):
        if paragraph.strip():para(paragraph)
    if title=='Declaration':
        for person in ['Sandesh Brahmane','Omkar Sonawane','Mahesh Patil','Sarvesh Mane']:
            p=para(person+'    Signature: __________________________');p.paragraph_format.space_before=Pt(10)
        para('Date: __________________    Place: __________________')
    elif title=='Acknowledgements':
        p=doc.add_paragraph(style='NEXOS Body');p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run('Sandesh Brahmane\nOmkar Sonawane\nMahesh Patil\nSarvesh Mane')

contents=[]
for title,instruction in [('Table of Contents','TOC \\t "NEXOS Front,1,NEXOS Chapter,1,NEXOS Section,2,NEXOS Appendix,1,NEXOS App Section,2" \\h \\z'),('List of Figures','TOC \\c "Figure" \\h \\z'),('List of Tables','TOC \\c "Table" \\h \\z')]:
    doc.add_paragraph(title,style='NEXOS Contents Title')
    marker=doc.add_paragraph(style='NEXOS Contents');contents.append((marker,instruction,title))

mainsection=doc.add_section(WD_SECTION_START.NEW_PAGE)
set_page_numbering(mainsection,'decimal')
for file in sorted((REPORT/'chapters').glob('[0-9][0-9]-*.tex')):
    parse(file.read_text(encoding='utf-8'))
heading('References',front=True)
for key,entry in re.findall(r'\\bibitem\{([^}]+)\}(.*?)(?=\\bibitem|\\end\{thebibliography\})',(REPORT/'references.tex').read_text(),re.S):
    p=para('['+BIB[key]+'] '+entry,'NEXOS Reference');p.paragraph_format.left_indent=Mm(8);p.paragraph_format.first_line_indent=Mm(-8)
appendix=True;chapter=0;secno=0;subno=0
parse((REPORT/'chapters'/'appendices.tex').read_text())

# Cached field contents show usable navigation immediately; Word refresh adds page numbers.
for marker,instruction,title in contents:
    entries=headings if title=='Table of Contents' else figures if title=='List of Figures' else tables
    field(marker,instruction)
    # Complex fields can encompass paragraphs; move separator/end around cached entries.
    endrun=marker._p[-1];marker._p.remove(endrun)
    anchor=marker._p
    for label,name,level in entries:
        p=OxmlElement('w:p');ppr=el('pPr');ppr.append(el('pStyle',val=doc.styles['NEXOS Contents'].style_id))
        if level:ppr.append(el('ind',left=250))
        p.append(ppr)
        link=el('hyperlink',anchor=name,history='1');r=el('r');t=el('t');t.text=label;r.append(t);link.append(r);p.append(link)
        anchor.addnext(p);anchor=p
    anchor.append(endrun)

settings=doc.settings.element
old=settings.find(qn('w:updateFields'))
if old is not None:settings.remove(old)
settings.append(el('updateFields',val='true'))
doc.core_properties.title='NEXOS Professional Project Report'
doc.core_properties.subject='Editable project report for review'
doc.core_properties.author='NEXOS Project Team'
doc.core_properties.keywords='Linux, AI agents, architecture, project report'
doc.core_properties.comments=''
OUTPUT.parent.mkdir(parents=True,exist_ok=True)
intermediate=QA/'generated.docx';doc.save(intermediate)

# Keep source parts that python-docx did not need to modify, including media.
with ZipFile(SOURCE) as original, ZipFile(intermediate) as generated:
    parts={n:original.read(n) for n in original.namelist()}
    parts.update({n:generated.read(n) for n in generated.namelist()})
    srcroot=etree.fromstring(original.read('word/document.xml'))
    outroot=etree.fromstring(parts['word/document.xml'])
    srcbody=srcroot.find('w:body',NS);outbody=outroot.find('w:body',NS)
    for i in range(74):outbody.replace(outbody[i],copy.deepcopy(srcbody[i]))
    parts['word/document.xml']=etree.tostring(outroot,xml_declaration=True,encoding='UTF-8',standalone=True)
    # Word resaves can change serialization; preserve original style definitions structurally.
    srcstyles=etree.fromstring(original.read('word/styles.xml'));outstyles=etree.fromstring(parts['word/styles.xml'])
    ids={s.get(qn('w:styleId')):s for s in outstyles.findall('w:style',NS)}
    for s in srcstyles.findall('w:style',NS):
        sid=s.get(qn('w:styleId'))
        if sid in ids:outstyles.replace(ids[sid],copy.deepcopy(s))
    parts['word/styles.xml']=etree.tostring(outstyles,xml_declaration=True,encoding='UTF-8',standalone=True)
    # Remove stale pagination statistics inherited from the shorter source report.
    app=etree.fromstring(parts['docProps/app.xml'])
    for child in list(app):
        if etree.QName(child).localname in ('Pages','Words','Characters','CharactersWithSpaces','Lines','Paragraphs'):
            app.remove(child)
    parts['docProps/app.xml']=etree.tostring(app,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(OUTPUT,'w',ZIP_DEFLATED) as z:
        for n,data in parts.items():z.writestr(n,data)

with ZipFile(OUTPUT) as z,ZipFile(SOURCE) as orig:
    final=etree.fromstring(z.read('word/document.xml'));fb=final.find('w:body',NS)
    def canon(e):return etree.tostring(e,method='c14n',exclusive=True)
    preserved=all(canon(fb[i])==canon(srcbody[i]) for i in range(74))
    assert preserved,'Protected front XML changed'
    newtext=' '.join(final.xpath('.//w:t/text()',namespaces=NS))
    unresolved=re.findall(r'\\(?:textbf|emph|cite|ref|begin|end|caption|section|chapter|url)\b',newtext)
    assert not unresolved,unresolved
    for n in orig.namelist():
        if n.startswith('word/media/'):assert z.read(n)==orig.read(n)
    evidence={'output':str(OUTPUT),'protected_front_elements_unchanged':preserved,
              'original_media_unchanged':True,'new_figures':len(figures),'captioned_tables':len(tables),
              'navigation_entries':len(headings),'references':len(BIB),
              'editable_tables':len(Document(OUTPUT).tables),
              'field_refresh_on_open':True,'word_render_verified':False,
              'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
              'output_sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest()}
    (QA/'verification.json').write_text(json.dumps(evidence,indent=2))
    print(json.dumps(evidence,indent=2))
